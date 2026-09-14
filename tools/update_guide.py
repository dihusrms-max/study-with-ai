from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


GUIDE = Path(r"C:\study-contests\Guide\나라장터_DACON_AI_경진대회_실전가이드.docx")


def insert_paragraph_before(anchor, text="", style=None, bold_prefix=None):
    paragraph = Paragraph(OxmlElement("w:p"), anchor._parent)
    anchor._p.addprevious(paragraph._p)
    if style:
        paragraph.style = style
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


doc = Document(GUIDE)
anchor = next(
    (p for p in doc.paragraphs if p.text.strip().startswith("15.")),
    None,
)
if anchor is None:
    raise RuntimeError("Section 15 heading was not found")

items = [
    ("16. V0 제출 정리", "Heading 1", None),
    (
        "V0의 목적은 점수 향상이 아니라 제출 구조와 실행 흐름을 빠르게 확인하는 것이다.",
        None,
        None,
    ),
    (
        "--mock 실행은 모델 없이 입력·출력 형식만 확인하는 로컬 테스트이므로 실제 리더보드 제출용이 아니다.",
        None,
        None,
    ),
    (
        "실제 V0 제출물은 submit.zip이며, 최상위에 script.py를 포함해야 한다. requirements.txt는 선택 사항이고, model/에는 RAG 인덱스와 설정 같은 정적 자산만 포함할 수 있다. 모델 가중치는 제출하지 않는다.",
        None,
        None,
    ),
    (
        "평가 서버는 data/와 output/을 제공한다. script.py는 평가 데이터에 대해 실제 고정 Gemma LLM을 호출하고 output/submission.csv를 생성해야 한다.",
        None,
        None,
    ),
    (
        "submission.csv는 id, v1~v24, e1~e24의 49개 열로 구성한다. v1~v24는 0 또는 1만 허용되며, id 누락·중복이 없어야 하고 UTF-8(BOM 없음)으로 저장한다.",
        None,
        None,
    ),
    (
        "제출 전 실제 LLM 호출, output/submission.csv 생성, 열 구조, 전체 행 수를 확인한다. 제출 횟수는 하루 1회이므로 mock 파일을 제출하지 않는다.",
        None,
        None,
    ),
]

for text, style, _ in reversed(items):
    insert_paragraph_before(anchor, text, style)

if not any(p.text.strip().startswith("17. 라벨링") for p in doc.paragraphs):
    doc.add_heading("17. 라벨링과 법령 해석 보완", level=1)
    doc.add_paragraph(
        "이 대회의 성능은 공고문과 배포된 법령 패키지를 비교하여 각 항목의 위반 여부를 정확히 판단하는 데서 시작한다. 단순 키워드 검색이 아니라 적용 법령, 요건, 예외 조건을 함께 확인한다."
    )
    doc.add_paragraph("항목별 판단 기준표는 다음 열로 관리한다.")
    doc.add_paragraph(
        "항목 | 관련 법령 | 위반 조건 | 예외 조건 | 근거 문장 기준",
        style="List Bullet",
    )
    doc.add_paragraph(
        "라벨은 v1~v24에 위반 여부를 기록하고, 위반이면 e1~e24에 제공 원문에서 그대로 인용한 근거를 기록한다. 부재 탐지 항목은 인용할 문장이 없을 수 있으므로 별도 기준을 둔다.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "초기 50~100건은 팀원이 일부 같은 샘플을 독립적으로 라벨링한다. 결과를 비교하여 불일치 사례를 합의하고, 확정한 기준으로 나머지 데이터를 분담한다.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "dev 데이터는 기준 확인과 검증에 사용한다. 라벨링 기준을 정한 뒤 일부 데이터를 검증용으로 남겨 기준 변경 전후의 차이를 확인한다.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "자가 라벨링을 사용하면 라벨 파일, 생성 프롬프트, 사용 모델, 생성 날짜와 버전을 함께 보관하여 재현 가능하게 관리한다.",
        style="List Bullet",
    )
    doc.add_paragraph("라벨링 완료 전 체크리스트")
    for text in [
        "[ ] v1~v24의 판단 기준과 예외 조건을 정리했다.",
        "[ ] 팀원 간 불일치 사례를 비교하고 합의했다.",
        "[ ] v 값은 0 또는 1만 사용했다.",
        "[ ] e 값은 제공 문서의 원문 부분 문자열로 기록했다.",
        "[ ] 라벨 파일과 생성 과정을 버전별로 보관했다.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

if not any(p.text.strip().startswith("18. 전체 작업 흐름") for p in doc.paragraphs):
    doc.add_heading("18. 전체 작업 흐름", level=1)
    doc.add_paragraph(
        "V0에서는 실행과 제출 구조를 확인하고, 01 노트북에서는 데이터와 라벨 구조를 확인한다. 이후 대표 사례를 분석하여 v 항목과 근거 문장을 이해하고, 그 판단을 Gemma가 재현하도록 프롬프트와 파이프라인을 개선한다."
    )
    doc.add_paragraph(
        "V0: 실행·제출 구조 확인 → 01 노트북: 데이터·라벨 구조 확인 → 대표 사례 분석 → 법령·예외 조건 학습 → Gemma 프롬프트 개선 → v1~v24 자동 예측 → dev 정답과 비교 → V1 제출"
    )
    doc.add_paragraph(
        "dev_labels.csv는 대회 평가용 정답으로 사용한다. 법령 검토는 라벨을 임의로 뒤집기 위한 것이 아니라, Gemma가 대회 기준과 같은 판단을 하도록 이해하기 위한 과정이다.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "사람이 샘플을 분석할 때는 관련 단서를 중심으로 검토할 수 있지만, 최종 제출 코드는 각 공고에 대해 v1~v24 전체 값을 출력해야 한다.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "2만 건 전체를 수동 라벨링하는 것은 필수가 아니다. 대표 샘플로 기준을 정리한 뒤 필요한 경우에만 자가 라벨링을 확대한다.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "PPS-DEV-01 사례는 특정 사례의 설명이며, 모든 공고에 같은 결론을 적용하지 않는다. 각 공고의 전체 문서와 항목별 기준을 확인한다.",
        style="List Bullet",
    )

if not any(p.text.strip().startswith("19. 노트북 실습 절차") for p in doc.paragraphs):
    doc.add_heading("19. 노트북 실습 절차", level=1)
    doc.add_paragraph(
        "01_data_inspection.ipynb는 다음 순서로 실행한다. 먼저 셀 1~5를 위에서부터 실행하여 ZIP, dev 데이터, 문서 구조와 dev_labels.csv를 준비한다. 커널을 다시 시작한 경우에도 이 순서를 반복한다."
    )
    doc.add_paragraph(
        "이후 대표 공고를 확인하고 summary_df를 만들어 항목별 대표 ID와 근거 문장을 정리한다.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "공고문 전체를 매번 출력하지 않는다. 대표 ID, 문서 종류, 위반 항목과 근거를 먼저 보고, 필요한 경우에만 원문을 텍스트 파일로 저장해 확인한다.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "첫 10건만으로 전체 경향을 판단하지 않는다. v1~v24 각각에서 위반 사례가 있는 대표 공고를 확인한다.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "dev_labels.csv의 정답을 확인할 때는 v1~v24를 위반 여부로 보고, v=1인 항목의 e 열에서 근거를 확인한다. v=0인 항목의 e 열은 비워 둔다.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "2만 건은 수동으로 모두 라벨링하지 않는다. 대표 샘플로 기준을 익힌 뒤 Gemma 자동 예측과 필요 시 자가 라벨링으로 확대한다.",
        style="List Bullet",
    )

doc.save(GUIDE)
print(GUIDE)
