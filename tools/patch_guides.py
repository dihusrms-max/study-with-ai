from pathlib import Path

from docx import Document


ROOT = Path(r"C:\study-contests")
GUIDE_DIR = ROOT / "Guide"


def replace_in_paragraphs(doc: Document, replacements: dict[str, str]) -> int:
    changed = 0
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        old = paragraph.text
        new = old
        for src, dst in sorted(replacements.items(), key=lambda item: len(item[0]), reverse=True):
            new = new.replace(src, dst)
        if new != old:
            paragraph.text = new
            changed += 1
    return changed


def add_special_section(path: Path) -> bool:
    doc = Document(path)
    if any(p.text.strip().startswith("20. 제출·운영 특이사항") for p in doc.paragraphs):
        return False

    doc.add_heading("20. 제출·운영 특이사항", level=1)
    doc.add_paragraph(
        "제출 기회는 채점 완료 시점이 아니라 제출 시점 기준으로 하루 1회만 주어진다. "
        "대기 중인 제출도 해당 날짜의 기회를 사용하므로, 제출 전 로컬 검증과 후보 버전 동결을 완료한다."
    )
    doc.add_paragraph(
        "공식 일정은 운영 상황에 따라 바뀔 수 있으므로 최종 제출 전 DACON 일정 페이지를 다시 확인한다. "
        "현재 기준 주요 일정은 팀 병합 2026년 9월 23일 23:59, 리더보드 제출 2026년 9월 29일 10:00, "
        "2차 평가 자료 제출 2026년 9월 30일 12:00부터 10월 5일 10:00까지다."
    )
    doc.add_paragraph(
        "프로젝트 기준 경로는 현재 저장소명인 koneps-dacon으로 통일한다. "
        "실험 코드는 src/koneps_dacon, 로컬 결과는 outputs, 제출 코드는 submit/script.py를 사용한다. "
        "가이드에 남아 있는 koneps-ai 또는 koneps-aicd 표기는 기존 문맥상 예시일 수 있으므로 실제 실행 경로와 혼동하지 않는다."
    )
    doc.add_paragraph(
        "설치 명령은 한 줄에 여러 명령을 붙여 실행하지 않는다. 각 명령을 별도 줄에 두고 한 단계씩 실행하며, "
        "실패 시 Linux → Driver → Docker → NVIDIA Container Toolkit → CUDA 컨테이너 → vLLM → Gemma 순서로 확인한다."
    )
    doc.add_paragraph(
        "최종 평가는 리더보드 점수만으로 끝나지 않으며, 근거 문구의 원문 일치성과 재현 자료가 별도 평가 대상이 될 수 있다. "
        "따라서 실험별 Prompt, 모델 설정, 실행 시간, FP/FN 분석과 evidence 검증 결과를 함께 보관한다."
    )
    doc.save(path)
    return True


def add_gemma_phase_section(path: Path) -> bool:
    doc = Document(path)
    if any(p.text.strip().startswith("21. Gemma 접근 전후 작업 구분") for p in doc.paragraphs):
        return False

    doc.add_heading("21. Gemma 접근 전후 작업 구분", level=1)
    doc.add_paragraph(
        "전체 작업은 Gemma 실행 환경을 확보하기 전과 후로 나누어 진행한다. "
        "Gemma 접근 전에는 받은 데이터와 배포 자료만으로 준비·분석하고, 접근 후에는 실제 모델 추론과 성능 비교를 수행한다."
    )
    doc.add_heading("21.1 Gemma 접근 전", level=2)
    doc.add_paragraph(
        "현재 받은 데이터만으로 dev.jsonl.gz와 dev_labels.csv의 구조를 확인하고, 항목표.json과 법령 패키지를 읽어 v1~v24의 정상·위반·예외 조건을 정리한다. "
        "대표 공고를 검토하고, 금액·비율·지역·업종·메타데이터 비교와 부재탐지처럼 규칙화할 수 있는 조건을 설계한다. "
        "또한 Prompt 초안, 실험 기록 양식, 제출 전 preflight, CSV 형식 검증을 준비한다."
    )
    doc.add_paragraph(
        "이 단계의 --mock 실행은 입력·출력 형식 확인용이다. 실제 Gemma가 아니므로 Macro F1이나 Prompt 성능 개선을 판단하는 용도로 사용하지 않는다."
    )
    doc.add_heading("21.2 Gemma 접근 후", level=2)
    doc.add_paragraph(
        "Gemma 실행 환경이 확보되면 공고 1건 추론으로 JSON 출력과 모델 설정을 확인한 뒤 dev 200건 전체를 실행한다. "
        "실제 Macro F1, 항목별 F1, FP/FN, JSON 오류율, 입력·출력 토큰, 전체 실행시간을 기록한다."
    )
    doc.add_paragraph(
        "이후 Prompt 개선, 규칙 기반 후처리, 필요한 항목에 한정한 RAG, 배치·속도 최적화를 한 번에 하나씩 비교한다. "
        "각 변경은 별도 버전으로 저장하고, 개선 결과를 확인한 뒤에만 다음 제출 후보로 승인한다."
    )
    doc.add_heading("21.3 단계별 완료 기준", level=2)
    doc.add_paragraph(
        "Gemma 접근 전 완료 기준은 데이터·라벨·항목 기준 확인, 규칙 후보 정리, Prompt 초안, preflight와 실험 기록 준비다."
    )
    doc.add_paragraph(
        "Gemma 접근 후 완료 기준은 공고 1건 추론 성공, dev 200건 평가 성공, 항목별 오류 분석표 작성, 실행시간 확인과 제출 후보 검증이다."
    )
    doc.save(path)
    return True


def dedupe_v0_section(path: Path) -> int:
    doc = Document(path)
    headings = [p for p in doc.paragraphs if p.text.strip() == "16. V0 제출 정리"]
    removed = 0
    for heading in headings[1:]:
        current = heading
        for _ in range(7):
            nxt = current._element.getnext()
            if nxt is None:
                break
            current._element.getparent().remove(current._element)
            removed += 1
            current = next((p for p in doc.paragraphs if p._element is nxt), None)
            if current is None:
                break
    doc.save(path)
    return removed


def add_operations_section(path: Path) -> bool:
    doc = Document(path)
    if any(p.text.strip().startswith("22. 운영 상태와 승인 절차") for p in doc.paragraphs):
        return False

    doc.add_heading("22. 운영 상태와 승인 절차", level=1)
    doc.add_paragraph(
        "현재 완료된 작업은 V0 실제 제출과 점수 확보, 데이터·dev 분포 분석, 우선 항목 사례 검토, "
        "판정 기준표 초안 작성, 가이드 보완이다. Gemma 서버 접근과 실제 dev 추론은 별도 요청 전까지 보류한다."
    )
    doc.add_paragraph(
        "사람이 법령 기준과 예외 조건을 검토하기 전에는 Prompt·Rule·RAG를 제출 코드에 반영하지 않는다. "
        "사람이 문서를 제공하면 관련 조항 추출, 24개 항목 매핑, 기준표 초안, 검토 대기 목록 순서로 정리한다."
    )
    doc.add_paragraph(
        "승인된 항목만 구현 후보로 전환한다. 변경 전후에는 dev 검증, 실행시간, FP/FN, evidence 원문 일치 여부를 기록하고, "
        "사람의 최종 승인을 받은 후보만 DACON에 제출한다."
    )
    doc.add_paragraph(
        "DACON 리더보드 점수, dev 로컬 점수, mock 형식 검증 결과는 서로 다른 지표이므로 실험 기록에서 구분한다."
    )
    doc.add_paragraph(
        "제출 전에는 실제 Gemma 호출, output/submission.csv 생성, 49개 열, 입력 행 수, id 유일성, v값 0·1, "
        "evidence 원문 부분문자열, 실행시간 제한을 확인한다."
    )
    doc.save(path)
    return True


def add_evaluation_section(path: Path) -> bool:
    doc = Document(path)
    if any(p.text.strip().startswith("23. 공식 평가 기준 반영") for p in doc.paragraphs):
        return False

    doc.add_heading("23. 공식 평가 기준 반영", level=1)
    doc.add_paragraph(
        "리더보드 점수는 전체 평가 데이터 1,853건에 대해 v1~v24의 위반 클래스 F1을 평균한 Macro F1이다. "
        "Public Score와 대회 종료 시점의 Private Score를 일반적인 데이터 분할 대회처럼 해석하지 않는다."
    )
    doc.add_paragraph(
        "1차 점수에는 v1~v24만 반영된다. e1~e24는 리더보드 점수에는 반영되지 않지만, 2차 평가에서 근거 문구 정성평가 자료로 사용된다."
    )
    doc.add_paragraph(
        "evidence는 공고문·첨부 문서·과업지시서·규격서·제안요청서에 실제 존재하는 부분문자열이어야 한다. "
        "요약·수정·생성한 문장과 나라장터 meta 값 자체는 evidence로 사용하지 않는다. "
        "v10·v11·v16·v18·v20은 부재탐지형 항목이므로 인용할 원문이 없고 evidence 정성평가 대상에서 제외된다."
    )
    doc.add_paragraph(
        "v24는 문서와 meta의 불일치를 판정하되, 불일치 근거는 meta 값이 아니라 문서에서 추출한 원문으로 남긴다. "
        "문서 필드 추출, meta 값 정규화, 비교, evidence 저장을 각각 기록한다."
    )
    doc.add_paragraph(
        "설치 오류와 제출 오류를 구분한다. ZIP 구조나 패키지 설치 실패는 설치 오류이고, script.py 실행 중 발생하는 오류는 제출 오류가 될 수 있다. "
        "패키지 설치 제한은 10분, script.py 전체 실행 제한은 2시간이며 모델 로드와 초기화 시간도 포함된다."
    )
    doc.add_paragraph(
        "requirements.txt에는 필요한 추가 패키지만 적고, 평가 서버에 고정 설치된 vllm·torch·transformers·xgrammar의 다른 버전을 설치하지 않는다. "
        "실행 중 인터넷 연결이나 모델 다운로드를 시도하지 않는다."
    )
    doc.add_paragraph(
        "2차 평가 재현을 위해 제출 ZIP, Prompt, 모델 설정, 모델 revision, vLLM 버전, 양자화 설정, max_model_len, seed, "
        "batch size, 실행 로그와 submission.csv를 버전별로 보관한다."
    )
    doc.add_paragraph(
        "제출 전에는 49개 열, 입력 행 수, id 유일성, v값 0·1, evidence 원문 일치, 설치 가능 여부와 전체 실행시간을 확인한다."
    )
    doc.save(path)
    return True


def main() -> None:
    replacements = {
        "koneps-aicd": "koneps-dacon",
        "koneps-ai": "koneps-dacon",
        "wsl --install -d Ubuntu-22.04wsl --set-default-version 2wsl --status":
            "wsl --install -d Ubuntu-22.04\nwsl --set-default-version 2\nwsl --status",
        "sudo apt updatesudo apt install": "sudo apt update\nsudo apt install",
        "uv inituv python pin 3.12.13": "uv init\nuv python pin 3.12.13",
        "cd ~/projects/koneps-daconuv run": "cd ~/projects/koneps-dacon\nuv run",
        "버전: v1.0": "버전: v1.4",
        "버전: v1.3": "버전: v1.4",
        "작성일: 2026년 9월 11일": "작성일: 2026년 9월 14일",
    }
    changed = 0
    for path in GUIDE_DIR.glob("*.docx"):
        doc = Document(path)
        changed += replace_in_paragraphs(doc, replacements)
        doc.save(path)
    practical = GUIDE_DIR / "나라장터_DACON_AI_경진대회_실전가이드.docx"
    removed = dedupe_v0_section(practical)
    added = add_special_section(practical)
    phases_added = add_gemma_phase_section(practical)
    operations_added = add_operations_section(practical)
    evaluation_added = add_evaluation_section(practical)
    print(f"paragraphs_changed={changed} v0_duplicates_removed={removed} special_section_added={added} gemma_phase_section_added={phases_added} operations_section_added={operations_added} evaluation_section_added={evaluation_added}")


if __name__ == "__main__":
    main()
