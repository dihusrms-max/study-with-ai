from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\study-with-ai\나라장터_DACON_개인작업환경_설치가이드.docx"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(table, color='D9D9D9'):
    tblPr = table._tbl.tblPr
    b = tblPr.first_child_found_in('w:tblBorders')
    if b is None:
        b = OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        el = b.find(qn(tag))
        if el is None: el = OxmlElement(tag); b.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), color)

def set_cell(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(9.5)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'; borders(t)
    for i,h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, True, 'FFFFFF'); shade(t.rows[0].cells[i], '274C77')
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            set_cell(cells[i], str(val));
            if ri%2==1: shade(cells[i], 'F3F6F9')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def code(doc, text):
    p=doc.add_paragraph(); p.style='No Spacing'; p.paragraph_format.left_indent=Inches(.2); p.paragraph_format.space_after=Pt(6)
    r=p.add_run(text); r.font.name='Consolas'; r._element.rPr.rFonts.set(qn('w:ascii'),'Consolas'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Consolas'); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string('333333')

doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕'); styles['Normal'].font.size=Pt(10)
for name,size in [('Title',22),('Heading 1',15),('Heading 2',12)]:
    styles[name].font.name='Aptos'; styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕'); styles[name].font.size=Pt(size); styles[name].font.color.rgb=RGBColor(0,0,0)

p=doc.add_paragraph(style='Title'); p.add_run('나라장터 DACON 개인 작업환경 설치 가이드')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
r=p.add_run('목표: 대회 서버에 접속하기 전에 개인 PC에서 분석·Baseline 개발을 재현 가능한 상태로 만든다.'); r.bold=True; r.font.size=Pt(11)
doc.add_paragraph('이 문서는 Windows 개인 PC 기준의 1단계 작업 환경만 다룹니다. 먼저 WSL2 Ubuntu 안에서 프로젝트를 실행하고, Python 패키지와 Jupyter가 정상 동작하는지 확인합니다. GPU 서버 접속, Docker, vLLM, Gemma 모델 실행은 개인 환경이 안정화된 뒤 별도 단계로 진행합니다.')

doc.add_heading('1. 이번 단계의 범위', level=1)
add_table(doc, ['구분','이번에 한다','이번에는 보류한다'], [
    ('개인 PC 기본 환경','WSL2 Ubuntu, VS Code, Git','대회 서버 SSH 접속'),
    ('Python 개발 환경','uv, Python 3.12.13, 가상환경','Conda 병행 설치'),
    ('분석 환경','pandas, numpy, scikit-learn, JupyterLab, pytest, ruff','CUDA Toolkit, PyTorch GPU'),
    ('프로젝트','koneps-ai 폴더와 기본 디렉터리','Docker, NVIDIA Toolkit, vLLM, Gemma'),
], [1.2,2.7,2.7])

doc.add_heading('2. 설치 순서', level=1)
doc.add_paragraph('아래 순서를 지키면 오류가 발생했을 때 어느 계층에서 문제가 생겼는지 확인하기 쉽습니다.')
doc.add_paragraph('Windows 기능 → WSL2 Ubuntu → VS Code 연동 → Git 확인 → uv와 Python → 프로젝트 생성 → 패키지 설치 → Jupyter 실행 → 최소 검증', style='Intense Quote')

doc.add_heading('3. Windows에 WSL2 Ubuntu 설치', level=1)
doc.add_paragraph('관리자 권한 PowerShell에서 실행합니다. 설치가 끝나면 재부팅하거나 안내에 따라 Ubuntu를 처음 실행해 사용자 계정을 만듭니다.')
code(doc, 'wsl --install -d Ubuntu-22.04\nwsl --set-default-version 2\nwsl --status')
doc.add_paragraph('Ubuntu 터미널에서 다음을 확인합니다.')
code(doc, 'cat /etc/os-release\nuname -r\nwhoami\npwd')
doc.add_paragraph('프로젝트 파일은 Windows 경로보다 WSL 내부의 홈 디렉터리(예: ~/projects)에 두는 것을 권장합니다. 파일 입출력이 많은 분석 작업에서 관리가 단순합니다.')

doc.add_heading('4. VS Code와 Git 확인', level=1)
doc.add_paragraph('Windows에 VS Code를 설치하고 확장 기능에서 Remote Development 또는 WSL을 설치합니다. Ubuntu 터미널에서 프로젝트 폴더로 이동한 뒤 `code .`으로 엽니다.')
code(doc, 'sudo apt update\nsudo apt install -y git git-lfs curl wget unzip zip build-essential tmux htop tree jq\ngit lfs install\ngit --version\ncode .')
doc.add_paragraph('Git 사용자 정보는 개인 계정으로 사용할 값이 정해졌을 때만 등록합니다. 대회 저장소나 원격 저장소 주소는 아직 추가하지 않습니다.')

doc.add_heading('5. uv와 Python 3.12.13 설치', level=1)
doc.add_paragraph('Ubuntu 터미널에서 실행합니다.')
code(doc, 'curl -LsSf https://astral.sh/uv/install.sh | sh\nexport PATH="$HOME/.local/bin:$PATH"\nuv --version\nuv python install 3.12.13\nuv python list')
doc.add_paragraph('새 터미널을 열 때도 uv가 인식되지 않으면 ~/.bashrc에 PATH가 등록되어 있는지 확인합니다. Python 버전은 프로젝트별로 고정해 환경 차이를 줄입니다.')

doc.add_heading('6. 개인 프로젝트 생성', level=1)
code(doc, 'mkdir -p ~/projects/koneps-ai\ncd ~/projects/koneps-ai\nuv init\nuv python pin 3.12.13')
doc.add_paragraph('기본 폴더를 생성합니다. 원본 데이터는 data/raw에 보관하고, 전처리 결과는 data/processed에 저장합니다. 원본 파일은 직접 수정하지 않습니다.')
code(doc, 'mkdir -p data/raw data/processed docs notebooks src/koneps_ai configs prompts\nmkdir -p artifacts/indexes artifacts/chunks experiments/errors outputs tests scripts submit\ntree -L 2')

doc.add_heading('7. 분석 패키지 설치와 검증', level=1)
code(doc, 'uv add pandas numpy scikit-learn pyyaml orjson tqdm\nuv add --dev jupyterlab ipykernel pytest ruff\nuv run python --version\nuv run python -c "import pandas, numpy, sklearn; print(\'Python environment OK\')"')
doc.add_paragraph('Python 버전이 3.12.13으로 출력되고 import 검증이 성공하면 기본 개발 환경은 정상입니다. 이후부터는 `python` 대신 `uv run python`, `pip` 대신 `uv add`를 사용합니다.')

doc.add_heading('8. JupyterLab 실행', level=1)
code(doc, 'cd ~/projects/koneps-ai\nuv run jupyter lab --no-browser --ip=127.0.0.1 --port=8888')
doc.add_paragraph('터미널에 표시되는 http://127.0.0.1:8888/?token=... 주소를 Windows 브라우저에서 엽니다. 8888 포트를 외부에 공개하지 않습니다. 종료할 때는 터미널에서 Ctrl+C를 누릅니다.')

doc.add_heading('9. 완료 체크리스트', level=1)
for item in ['WSL2 Ubuntu 22.04 실행 성공','VS Code에서 WSL 프로젝트 폴더 열기 성공','Git / curl / tmux 등 기본 명령 실행 성공','uv 설치 및 버전 출력 성공','Python 3.12.13 pin 성공','pandas / numpy / scikit-learn import 성공','koneps-ai 기본 폴더 생성','JupyterLab 실행 및 브라우저 접속 성공','원본 데이터와 결과 데이터 저장 위치 구분']:
    doc.add_paragraph('☐ ' + item)

doc.add_heading('10. 문제가 생겼을 때 확인 순서', level=1)
add_table(doc, ['증상','먼저 확인할 것','조치'], [
    ('wsl 명령 실패','Windows 재부팅 및 WSL 상태','wsl --status, wsl -l -v'),
    ('code 명령 없음','VS Code와 WSL 확장 설치','VS Code에서 폴더를 직접 열기'),
    ('uv 명령 없음','~/.local/bin PATH','새 터미널을 열고 PATH 재확인'),
    ('Python 버전 다름','.python-version 파일','uv python pin 3.12.13 재실행'),
    ('import 실패','현재 폴더와 uv 환경','uv add로 패키지 설치 후 uv run 사용'),
    ('Jupyter 접속 실패','실행 터미널의 token과 포트','Jupyter를 재실행하고 127.0.0.1 주소 사용'),
], [1.4,2.5,2.7])

doc.add_heading('11. 다음 단계로 미룰 항목', level=1)
doc.add_paragraph('개인 V0 환경이 완료된 뒤에만 GPU/V1을 검토합니다. NVIDIA 드라이버, Docker Engine, NVIDIA Container Toolkit, CUDA 컨테이너, vLLM, Gemma 모델은 개인 PC의 GPU 사양과 목적을 확인한 후 별도 설치 가이드로 분리합니다. 대회 서버 접근과 제출 설정은 그 이후에도 개인 환경과 섞지 않습니다.')
doc.add_paragraph('기준 버전: 원본 가이드의 Python 3.12.13 및 Ubuntu 22.04 기준 / 정리일: 2026-09-11', style='Caption')

doc.save(OUT)
print(OUT)
